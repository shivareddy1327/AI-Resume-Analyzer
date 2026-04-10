"""
AI Resume Analyzer using Google Gemini for Smart Resume AI
"""
import os
import io
import json
import re
import streamlit as st
from datetime import datetime


class AIResumeAnalyzer:
    def __init__(self):
        self.gemini_api_key = os.getenv("GEMINI_API_KEY", "")
        if not self.gemini_api_key:
            # Try Streamlit secrets
            try:
                self.gemini_api_key = st.secrets.get("GEMINI_API_KEY", "")
            except Exception:
                pass

    def extract_text_from_pdf(self, file) -> str:
        """Extract text from PDF file"""
        try:
            import pdfplumber
            with pdfplumber.open(file) as pdf:
                text = ""
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            return text
        except Exception:
            pass

        try:
            from pypdf import PdfReader
            reader = PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except Exception as e:
            raise Exception(f"Could not extract text from PDF: {str(e)}")

    def extract_text_from_docx(self, file) -> str:
        """Extract text from DOCX file"""
        try:
            from docx import Document
            doc = Document(file)
            text = "\n".join([para.text for para in doc.paragraphs])
            return text
        except Exception as e:
            raise Exception(f"Could not extract text from DOCX: {str(e)}")

    def analyze_resume_with_gemini(self, resume_text: str, job_role: str = "", job_description: str = "") -> dict:
        """Analyze resume using Google Gemini API"""
        if not self.gemini_api_key:
            return {
                "error": "Gemini API key not configured. Please add GEMINI_API_KEY to your environment or Streamlit secrets.",
                "analysis": "",
                "resume_score": 0,
                "ats_score": 0,
                "model_used": "Google Gemini"
            }

        try:
            import google.generativeai as genai
            genai.configure(api_key=self.gemini_api_key)
            model = genai.GenerativeModel("gemini-1.5-flash")

            if job_description:
                prompt = f"""You are an expert resume reviewer and career coach. Analyze the following resume for the role of "{job_role}" and compare it against the provided job description.

RESUME:
{resume_text}

JOB DESCRIPTION:
{job_description}

Provide a comprehensive analysis with the following sections:

## Overall Assessment
Brief overall evaluation of the resume quality and fit for the role.

## Professional Profile Analysis
Analysis of candidate's background, experience level, and professional trajectory.

## Skills Analysis
- Skills present that match the job description
- Critical missing skills
- Recommended skills to add

## Experience Analysis
Evaluation of work experience relevance, impact, and presentation.

## Education Analysis
Assessment of educational background relevance.

## Key Strengths
Top 3-5 strengths of this resume.

## Areas for Improvement
Top 3-5 specific improvements needed.

## ATS Optimization Assessment
How well optimized this resume is for ATS systems.

## Job Match Analysis
How well the resume matches the specific job description.

## Recommended Courses
3-5 specific courses to improve candidacy.

## Resume Score
Overall score: X/100
ATS Score: Y/100
Job Match Score: Z/100

Provide actionable, specific feedback. Be honest but constructive."""

            else:
                prompt = f"""You are an expert resume reviewer and career coach. Analyze the following resume for the role of "{job_role}".

RESUME:
{resume_text}

Provide a comprehensive analysis with the following sections:

## Overall Assessment
Brief overall evaluation of the resume.

## Professional Profile Analysis
Analysis of candidate's background and experience level.

## Skills Analysis
- Present relevant skills
- Missing important skills for {job_role}

## Experience Analysis
Evaluation of work experience.

## Education Analysis
Assessment of educational background.

## Key Strengths
Top 3-5 strengths.

## Areas for Improvement
Top 3-5 specific improvements.

## ATS Optimization Assessment
ATS optimization evaluation.

## Role Alignment Analysis
How well this resume aligns with {job_role} requirements.

## Recommended Courses
3-5 specific courses.

## Resume Score
Overall score: X/100
ATS Score: Y/100

Be specific and actionable."""

            response = model.generate_content(prompt)
            analysis_text = response.text

            # Parse scores from response
            resume_score = self._extract_score(analysis_text, "overall score")
            ats_score = self._extract_score(analysis_text, "ats score")
            job_match_score = self._extract_score(analysis_text, "job match score") if job_description else 0

            return {
                "analysis": analysis_text,
                "resume_score": resume_score,
                "ats_score": ats_score,
                "job_match_score": job_match_score,
                "model_used": "Google Gemini",
                "strengths": self._extract_section(analysis_text, "Key Strengths"),
                "weaknesses": self._extract_section(analysis_text, "Areas for Improvement"),
            }

        except Exception as e:
            return {
                "error": str(e),
                "analysis": "",
                "resume_score": 0,
                "ats_score": 0,
                "model_used": "Google Gemini"
            }

    def _extract_score(self, text: str, score_type: str) -> int:
        """Extract a score from analysis text"""
        pattern = rf'{score_type}[:\s]*(\d+)[/\s]*100'
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return min(int(match.group(1)), 100)
        # Fallback: look for any X/100 pattern near the score_type
        lines = text.lower().split('\n')
        for i, line in enumerate(lines):
            if score_type.lower() in line:
                score_match = re.search(r'(\d+)/100', line)
                if score_match:
                    return min(int(score_match.group(1)), 100)
        return 65  # Default score

    def _extract_section(self, text: str, section_name: str) -> list:
        """Extract bullet points from a named section"""
        pattern = rf'## {section_name}\n(.*?)(?=\n##|\Z)'
        match = re.search(pattern, text, re.DOTALL | re.IGNORECASE)
        if match:
            section_text = match.group(1)
            items = re.findall(r'[-•*]\s+(.+)', section_text)
            return [item.strip() for item in items]
        return []

    def generate_pdf_report(self, analysis_result: dict, candidate_name: str = "Candidate", job_role: str = "") -> bytes:
        """Generate a PDF report of the analysis"""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.colors import HexColor, white, black
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib import colors
            from reportlab.lib.units import inch

            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=letter,
                                    rightMargin=0.75 * inch, leftMargin=0.75 * inch,
                                    topMargin=0.75 * inch, bottomMargin=0.75 * inch)

            styles = getSampleStyleSheet()
            story = []

            # Title
            title_style = ParagraphStyle('Title', parent=styles['Title'],
                                         fontSize=20, textColor=HexColor('#1a1a2e'),
                                         spaceAfter=12)
            story.append(Paragraph("AI Resume Analysis Report", title_style))
            story.append(Spacer(1, 0.2 * inch))

            # Meta info
            meta_style = ParagraphStyle('Meta', parent=styles['Normal'],
                                        fontSize=11, textColor=HexColor('#444444'))
            story.append(Paragraph(f"<b>Candidate:</b> {candidate_name}", meta_style))
            story.append(Paragraph(f"<b>Role:</b> {job_role or 'Not specified'}", meta_style))
            story.append(Paragraph(f"<b>Date:</b> {datetime.now().strftime('%B %d, %Y')}", meta_style))
            story.append(Paragraph(f"<b>Model:</b> {analysis_result.get('model_used', 'Google Gemini')}", meta_style))
            story.append(Spacer(1, 0.3 * inch))

            # Scores table
            score = analysis_result.get('score', 0)
            ats_score = analysis_result.get('ats_score', 0)
            scores_data = [
                ['Metric', 'Score', 'Status'],
                ['Overall Resume Score', f"{score}/100",
                 'Excellent' if score >= 80 else 'Good' if score >= 60 else 'Needs Improvement'],
                ['ATS Optimization Score', f"{ats_score}/100",
                 'Excellent' if ats_score >= 80 else 'Good' if ats_score >= 60 else 'Needs Improvement'],
            ]

            table = Table(scores_data, colWidths=[3 * inch, 1.5 * inch, 2 * inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), HexColor('#1a1a2e')),
                ('TEXTCOLOR', (0, 0), (-1, 0), white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 11),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [HexColor('#f8f9fa'), white]),
                ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#dddddd')),
                ('PADDING', (0, 0), (-1, -1), 8),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.3 * inch))

            # Analysis text
            body_style = ParagraphStyle('Body', parent=styles['Normal'],
                                        fontSize=10, leading=14,
                                        textColor=HexColor('#333333'))
            heading_style = ParagraphStyle('Heading', parent=styles['Heading2'],
                                           fontSize=13, textColor=HexColor('#1a1a2e'),
                                           spaceBefore=12, spaceAfter=6)

            full_response = analysis_result.get('full_response', '')
            if full_response:
                for line in full_response.split('\n'):
                    if line.startswith('## '):
                        story.append(Paragraph(line.replace('## ', ''), heading_style))
                    elif line.strip().startswith('- ') or line.strip().startswith('* '):
                        story.append(Paragraph(f"• {line.strip()[2:]}", body_style))
                    elif line.strip():
                        story.append(Paragraph(line.strip(), body_style))

            doc.build(story)
            buffer.seek(0)
            return buffer.read()

        except ImportError:
            # Fallback: generate simple text-based PDF using fpdf2
            try:
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Helvetica", size=16, style='B')
                pdf.cell(0, 10, "AI Resume Analysis Report", ln=True, align='C')
                pdf.set_font("Helvetica", size=12)
                pdf.cell(0, 8, f"Candidate: {candidate_name}", ln=True)
                pdf.cell(0, 8, f"Role: {job_role}", ln=True)
                pdf.cell(0, 8, f"Score: {analysis_result.get('score', 0)}/100", ln=True)
                pdf.cell(0, 8, f"ATS Score: {analysis_result.get('ats_score', 0)}/100", ln=True)
                pdf.ln(5)
                pdf.set_font("Helvetica", size=10)
                full_text = analysis_result.get('full_response', 'No analysis available.')
                for line in full_text.split('\n'):
                    if line.strip():
                        pdf.multi_cell(0, 6, line[:200])
                return pdf.output()
            except Exception:
                return b""
        except Exception:
            return b""
